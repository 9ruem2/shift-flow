import io
import os
import json
import re
from datetime import date, datetime, timedelta
from typing import Tuple, List, Dict, Any, Optional
import openpyxl
from openpyxl.utils import get_column_letter

from app.core.exceptions import MissingWorkdayColumnError, FileParsingError
from app.services.week_calculator import WeekCalculator
from app.schemas.schedule import DriverSubstitutionRule, SubstitutionRecord

CONTRACT_DRIVERS_JSON_PATH = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "data",
    "contract_drivers.json"
)

CAMP_ROUTES_JSON_PATH = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "data",
    "camp_routes.json"
)

def load_contract_drivers_map() -> Dict[str, str]:
    """용차 기사 JSON 마스터 데이터 로드 (이름 -> ID 딕셔너리)"""
    if os.path.exists(CONTRACT_DRIVERS_JSON_PATH):
        try:
            with open(CONTRACT_DRIVERS_JSON_PATH, "r", encoding="utf-8") as f:
                data = json.load(f)
                return {item["name"].strip(): item["id"].strip() for item in data if "name" in item and "id" in item}
        except Exception:
            pass
    return {
        "이미복": "wlfjddp1",
        "윤주영": "rose2411",
        "지요셉": "j001919",
        "최수빈": "tomorrow1004",
        "배정한": "bjh7823",
        "오지훈": "daegook",
        "강민규": "mindalgod",
        "권광훈": "wlfjddp"
    }

def load_camp_routes_map() -> Dict[str, List[str]]:
    """캠프별 라우트 목록 JSON 로드"""
    if os.path.exists(CAMP_ROUTES_JSON_PATH):
        try:
            with open(CAMP_ROUTES_JSON_PATH, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            pass
    return {
        "구리2": ["905C,905D", "905A,905B", "005C,005D", "005A,005B", "003A,003B,003C,003D", "002A,002B,002C,002D", "905B,905C", "905A,905D", "002C,002D", "002A,002B"],
        "남양주1": ["211C,211D"],
        "남양주2": ["519C,519D", "518C,518D", "518A,518B", "517C,517D", "517B", "517A", "515C,515D"],
        "남양주3": ["905C,905D", "905A,905B", "903A,903B,903C,903D", "808B,808C,808D"],
        "남양주4": ["605D", "605C", "508A,508B", "505C,505D", "505A,505B", "504C", "504A,504B,504D", "502C,502D", "504C,504D", "504A,504B"]
    }

class ExcelShifter:
    """
    openpyxl 기반 날짜 및 용차 기사 치환 엔진
    1. 업무일 우선 탐색 및 대상 요일 매핑 (기존 서식/스타일/빈 셀 100% 보존)
    2. 변경된 새 업무일 기준 캠프/라우트 또는 기사명 탐색 후 용차 기사 정보로 치환
    """

    WORKDAY_HEADER_CANDIDATES = ["업무일", "업무일자", "업무_일자", "일자", "날짜", "WORK_DATE", "WORKDATE"]
    CAMP_HEADER_CANDIDATES = ["캠프", "캠프명", "소속캠프", "소속_캠프", "소속", "CAMP", "CAMP_NAME", "배송캠프"]
    ROUTE_HEADER_CANDIDATES = ["라우트", "라우트명", "소속라우트", "노선", "코스", "배송구역", "ROUTE", "ROUTE_NAME", "ROUTE_ID"]
    DRIVER_NAME_CANDIDATES = ["기사", "기사이름", "기사명", "성명", "기사_이름", "기사_성명", "기사_명", "DRIVER_NAME", "DRIVER", "NAME", "위탁기사", "배송기사"]
    DRIVER_ID_CANDIDATES = ["아이디", "기사아이디", "기사ID", "사번", "기사_아이디", "기사_ID", "ID", "DRIVER_ID", "EMPID", "EMP_ID", "위탁자ID", "위탁ID"]

    @classmethod
    def match_camp(cls, expected: Optional[str], actual: Optional[str]) -> bool:
        """캠프명 비교 (공백/대소문자 무시 및 부분 일치)"""
        if not expected or not expected.strip():
            return True
        if not actual:
            return False
        exp = expected.strip().replace(" ", "").lower()
        act = actual.strip().replace(" ", "").lower()
        return exp == act or exp in act or act in exp

    @classmethod
    def match_route(cls, expected: Optional[str], actual: Optional[str]) -> bool:
        """라우트명 비교 (공백/대소문자 무시, 토큰 집합 일치 등)"""
        if not expected or not expected.strip():
            return True
        if not actual:
            return False
        exp_clean = expected.strip().replace(" ", "").upper()
        act_clean = actual.strip().replace(" ", "").upper()
        if not act_clean:
            return False
        if exp_clean == act_clean or exp_clean in act_clean or act_clean in exp_clean:
            return True
        exp_tokens = set(re.split(r'[,/|\s]+', exp_clean)) - {""}
        act_tokens = set(re.split(r'[,/|\s]+', act_clean)) - {""}
        if exp_tokens and act_tokens:
            if exp_tokens == act_tokens or exp_tokens.issubset(act_tokens) or act_tokens.issubset(exp_tokens):
                return True
        return False

    @classmethod
    def parse_date_value(cls, val: Any) -> Optional[date]:
        """셀 값에서 date 객체 추출"""
        if val is None:
            return None
        if isinstance(val, datetime):
            return val.date()
        if isinstance(val, date):
            return val
        if isinstance(val, (int, float)):
            try:
                dt = openpyxl.utils.datetime.from_excel(val)
                return dt.date()
            except Exception:
                return None
        if isinstance(val, str):
            val_clean = val.strip()
            match = re.search(r'(\d{4})[-./](\d{1,2})[-./](\d{1,2})', val_clean)
            if match:
                y, m, d = map(int, match.groups())
                try:
                    return date(y, m, d)
                except ValueError:
                    return None
        return None

    @classmethod
    def format_date_like_original(cls, original_val: Any, new_dt: date) -> Any:
        """원본 값의 타입 및 문자열 포맷에 맞추어 변환된 날짜 반환"""
        if isinstance(original_val, datetime):
            return datetime(new_dt.year, new_dt.month, new_dt.day, original_val.hour, original_val.minute, original_val.second)
        if isinstance(original_val, date):
            return new_dt
        if isinstance(original_val, str):
            val_clean = original_val.strip()
            if "." in val_clean:
                return new_dt.strftime("%Y.%m.%d")
            elif "/" in val_clean:
                return new_dt.strftime("%Y/%m/%d")
            else:
                return new_dt.strftime("%Y-%m-%d")
        return datetime(new_dt.year, new_dt.month, new_dt.day)

    @classmethod
    def find_column_by_candidates(cls, ws, candidates: List[str], header_row: Optional[int] = None) -> Tuple[Optional[int], Optional[int], Optional[str]]:
        """
        주어진 후보 헤더명으로 컬럼 및 행 위치 탐색
        반환: (header_row_idx, col_idx, matched_header_name) (1-indexed)
        """
        start_r = header_row if header_row is not None else 1
        end_r = header_row if header_row is not None else min(15, ws.max_row or 15)
        max_search_col = min(50, ws.max_column or 50)

        for r in range(start_r, end_r + 1):
            for c in range(1, max_search_col + 1):
                val = ws.cell(row=r, column=c).value
                if val and isinstance(val, str):
                    clean_val = val.strip().replace(" ", "")
                    for candidate in candidates:
                        cand_clean = candidate.replace(" ", "")
                        if cand_clean == clean_val or cand_clean in clean_val or clean_val in cand_clean:
                            return r, c, val.strip()

        return None, None, None

    @classmethod
    def find_workday_column(cls, ws) -> Tuple[Optional[int], Optional[int], Optional[str]]:
        """
        워크시트에서 '업무일' 헤더 컬럼 및 행 위치 탐색
        반환: (header_row_idx, col_idx, matched_header_name) (1-indexed)
        """
        r, c, name = cls.find_column_by_candidates(ws, cls.WORKDAY_HEADER_CANDIDATES)
        if c is not None:
            return r, c, name

        max_search_row = min(15, ws.max_row or 15)
        max_search_col = min(50, ws.max_column or 50)

        date_counts_per_col = {}
        for col_idx in range(1, max_search_col + 1):
            count = 0
            for row_idx in range(1, max_search_row + 15):
                val = ws.cell(row=row_idx, column=col_idx).value
                if cls.parse_date_value(val) is not None:
                    count += 1
            if count > 0:
                date_counts_per_col[col_idx] = count

        if date_counts_per_col:
            best_col = max(date_counts_per_col, key=date_counts_per_col.get)
            for row_idx in range(1, max_search_row + 1):
                if cls.parse_date_value(ws.cell(row=row_idx, column=best_col).value) is not None:
                    header_r = max(1, row_idx - 1)
                    header_name = str(ws.cell(row=header_r, column=best_col).value or "업무일(추정)")
                    return header_r, best_col, header_name

        return None, None, None

    @classmethod
    def transform_schedule_for_week(
        cls,
        file_bytes: bytes,
        target_week_info: Dict[str, Any],
        substitution_rules: Optional[List[DriverSubstitutionRule]] = None
    ) -> Tuple[bytes, List[Dict[str, Any]], Dict[str, Any], List[SubstitutionRecord]]:
        """
        1. 특정 주차(일~토)에 맞추어 엑셀 파일의 업무일 날짜를 1:1 치환
        2. 치환된 새 업무일 및 캠프/라우트(또는 기존 기사명) 기준으로 용차 기사명 및 ID 치환
        """
        try:
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=False)
        except Exception as e:
            raise FileParsingError(f"엑셀 파일을 열 수 없습니다: {str(e)}")

        ws = wb.active
        if ws is None:
            raise FileParsingError("유효한 워크시트가 존재하지 않습니다.")

        header_row, workday_col, header_name = cls.find_workday_column(ws)
        if workday_col is None:
            raise MissingWorkdayColumnError("엑셀 파일 내에서 '업무일' 컬럼을 찾을 수 없습니다.")

        # 기사명, 아이디, 캠프, 라우트 컬럼 탐색
        _, driver_name_col, driver_name_header = cls.find_column_by_candidates(
            ws, cls.DRIVER_NAME_CANDIDATES, header_row=header_row
        )
        _, driver_id_col, driver_id_header = cls.find_column_by_candidates(
            ws, cls.DRIVER_ID_CANDIDATES, header_row=header_row
        )
        _, camp_col, camp_header = cls.find_column_by_candidates(
            ws, cls.CAMP_HEADER_CANDIDATES, header_row=header_row
        )
        _, route_col, route_header = cls.find_column_by_candidates(
            ws, cls.ROUTE_HEADER_CANDIDATES, header_row=header_row
        )

        target_sun = datetime.strptime(target_week_info["start_date"], "%Y-%m-%d").date()
        contract_drivers_map = load_contract_drivers_map()

        original_date_entries = []
        for r in range(header_row + 1, ws.max_row + 1):
            cell = ws.cell(row=r, column=workday_col)
            parsed_dt = cls.parse_date_value(cell.value)
            if parsed_dt is not None:
                original_date_entries.append({
                    "row": r,
                    "cell": cell,
                    "date": parsed_dt,
                    "original_val": cell.value,
                    "weekday": WeekCalculator.get_weekday_index(parsed_dt),
                    "day_name": WeekCalculator.get_korean_weekday_name(parsed_dt)
                })

        if not original_date_entries:
            raise MissingWorkdayColumnError(f"'{header_name}' 컬럼에 유효한 날짜 데이터가 없습니다.")

        # -------------------------------------------------------------
        # 1단계: 업무일 우선 치환 (새 주차 요일 매핑)
        # -------------------------------------------------------------
        mapping_records = []
        row_to_new_date: Dict[int, str] = {} # row_index -> YYYY-MM-DD

        for entry in original_date_entries:
            r = entry["row"]
            cell = entry["cell"]
            orig_dt = entry["date"]
            weekday_idx = entry["weekday"]

            # 대상 주차의 동일 요일 날짜로 치환
            new_date = target_sun + timedelta(days=weekday_idx)
            new_val = cls.format_date_like_original(entry["original_val"], new_date)
            cell.value = new_val

            new_date_str = new_date.strftime("%Y-%m-%d")
            row_to_new_date[r] = new_date_str

            mapping_records.append({
                "row_index": r,
                "original_date": orig_dt.strftime("%Y-%m-%d"),
                "transformed_date": new_date_str,
                "day_name": entry["day_name"],
                "target_week": target_week_info["formatted_range"]
            })

        # -------------------------------------------------------------
        # 2단계: 변경된 새 업무일 기준 용차 기사 및 ID 치환
        # -------------------------------------------------------------
        applied_substitutions: List[SubstitutionRecord] = []

        if substitution_rules and len(substitution_rules) > 0 and driver_name_col is not None:
            active_rules = []
            for rule in substitution_rules:
                t_date = (rule.target_date or "").strip()
                camp_val = (rule.camp or "").strip()
                route_val = (rule.route or "").strip()
                orig_name = (rule.original_driver_name or "").strip()
                new_name = (rule.new_driver_name or "").strip()
                new_id = rule.new_driver_id.strip() if rule.new_driver_id else contract_drivers_map.get(new_name, "")
                if t_date and new_name and (route_val or camp_val or orig_name):
                    active_rules.append({
                        "target_date": t_date,
                        "camp": camp_val,
                        "route": route_val,
                        "original_name": orig_name,
                        "new_name": new_name,
                        "new_id": new_id
                    })

            for r, new_workday_str in row_to_new_date.items():
                name_cell = ws.cell(row=r, column=driver_name_col)
                curr_driver_name = str(name_cell.value or "").strip()

                curr_camp = str(ws.cell(row=r, column=camp_col).value or "").strip() if camp_col else ""
                curr_route = str(ws.cell(row=r, column=route_col).value or "").strip() if route_col else ""

                for rule in active_rules:
                    # 1) 치환된 새 업무일 일치 여부 확인
                    if new_workday_str != rule["target_date"]:
                        continue

                    matched = False

                    # 2-A) 캠프 및 라우트 조건이 주어진 경우 매칭
                    if rule["route"] or rule["camp"]:
                        camp_ok = cls.match_camp(rule["camp"], curr_camp) if camp_col else True
                        route_ok = False
                        if route_col:
                            route_ok = cls.match_route(rule["route"], curr_route)
                        else:
                            # 라우트 컬럼이 별도로 없는 경우 전체 행 탐색
                            for c_idx in range(1, min(ws.max_column + 1, 30)):
                                c_val = str(ws.cell(row=r, column=c_idx).value or "")
                                if cls.match_route(rule["route"], c_val):
                                    route_ok = True
                                    break

                        if camp_ok and route_ok:
                            matched = True

                    # 2-B) 기존 기사명 조건이 주어진 경우 매칭
                    if not matched and rule["original_name"]:
                        if curr_driver_name == rule["original_name"] or rule["original_name"] in curr_driver_name:
                            matched = True

                    if matched:
                        # 기사 이름 변경
                        name_cell.value = rule["new_name"]

                        # 기사 ID 컬럼이 존재하면 ID도 변경
                        if driver_id_col is not None:
                            id_cell = ws.cell(row=r, column=driver_id_col)
                            id_cell.value = rule["new_id"]

                        applied_substitutions.append(
                            SubstitutionRecord(
                                target_date=new_workday_str,
                                camp=rule["camp"] or curr_camp or None,
                                route=rule["route"] or curr_route or None,
                                original_driver_name=curr_driver_name,
                                new_driver_name=rule["new_name"],
                                new_driver_id=rule["new_id"],
                                row_index=r,
                                applied=True
                            )
                        )
                        break

        out_stream = io.BytesIO()
        wb.save(out_stream)
        out_bytes = out_stream.getvalue()

        meta = {
            "workday_column_name": header_name,
            "workday_column_index": workday_col,
            "driver_name_column_name": driver_name_header,
            "driver_name_column_index": driver_name_col,
            "driver_id_column_name": driver_id_header,
            "driver_id_column_index": driver_id_col,
            "header_row": header_row,
            "total_rows": ws.max_row,
            "total_cols": ws.max_column,
            "workday_count": len(original_date_entries),
            "target_week_info": target_week_info
        }

        return out_bytes, mapping_records, meta, applied_substitutions

